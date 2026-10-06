import openpyxl,json,time
out={}
for fn in ["USD _ Регистрация.xlsx","USD _ Результаты.xlsx"]:
    wb=openpyxl.load_workbook(fn,data_only=True,read_only=True)
    fs={}
    for ws in wb.worksheets:
        it=ws.iter_rows(values_only=True); first=[]
        for i,row in enumerate(it,1):
            if i<=5:first.append(list(row))
            else:break
        fs[ws.title]={"max_row":ws.max_row,"max_col":ws.max_column,"sample":first}
    out[fn]=fs
print(json.dumps(out,ensure_ascii=False,separators=(",",":")),flush=True)
time.sleep(3600)
