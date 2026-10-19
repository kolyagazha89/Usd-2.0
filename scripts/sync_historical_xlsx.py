#!/usr/bin/env python3
"""Synchronize the checked-in historical JSON from the ACTUAL XLSX workbooks.

Usage:
    python scripts/sync_historical_xlsx.py --write
    python scripts/sync_historical_xlsx.py --check

Do not correct people's names or remove result rows in this conversion step.
Normalization is applied later by backend/src/imports/global-historical-plan.ts.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as error:
    raise SystemExit("Install converter dependency: python -m pip install openpyxl==3.1.5") from error

ROOT = Path(__file__).resolve().parents[1]
SNAPSHOT = ROOT / "backend/data/historical-xlsx.json"
WORKBOOKS = ("USD _ Регистрация.xlsx", "USD _ Результаты.xlsx")


def clean_cell(value):
    if value is None:
        return None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("Non-finite numeric cell value in XLSX")
        if value.is_integer():
            return int(value)
    if isinstance(value, (int, float, str, bool)):
        return value
    raise ValueError(f"Unsupported XLSX cell type: {type(value).__name__}")


def read_workbook(path: Path):
    book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        result = {}
        for sheet in book:
            # Retain ALL worksheet rows and columns, including the padded rows
            # currently present in the archived workbook. No data cleanup here.
            result[sheet.title] = [
                [clean_cell(cell) for cell in values]
                for values in sheet.iter_rows(values_only=True)
            ]
        return result
    finally:
        book.close()


def count_results(sheets: dict) -> int:
    """Count result entries using the same carried-class rule as the backend."""
    total = 0
    for name, rows in sheets.items():
        if not rows:
            continue
        headers = [str(cell) if cell is not None else "" for cell in rows[0]]
        try:
            person_col = headers.index("ФИО")
            class_col = headers.index("Класс")
        except ValueError:
            raise ValueError(f"Missing Results header in worksheet {name!r}")
        current_class = ""
        for values in rows[1:]:
            value = values[class_col] if class_col < len(values) else None
            if value is not None and str(value).strip():
                current_class = str(value).strip()
            person = values[person_col] if person_col < len(values) else None
            if current_class and person is not None and str(person).strip():
                total += 1
    return total


def create_snapshot():
    previous = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    data = {}
    workbooks = []
    for filename in WORKBOOKS:
        path = ROOT / filename
        if not path.is_file():
            raise FileNotFoundError(f"Historical source workbook missing: {path}")
        payload = path.read_bytes()
        if not payload.startswith(b"PK"):
            raise ValueError(f"Not an XLSX ZIP workbook: {path}")
        workbooks.append({
            "file": filename,
            "suppliedFile": filename,
            "sha256": hashlib.sha256(payload).hexdigest(),
        })
        data[filename] = read_workbook(path)
    historical_count = count_results(data[WORKBOOKS[1]])
    if historical_count < 1:
        raise ValueError("Refusing to publish an empty Results archive")
    former_count = count_results(previous[WORKBOOKS[1]])
    if historical_count < 500:
        raise ValueError(f"Results count unexpectedly low: {historical_count}; previous: {former_count}")
    changed_source = any(
        next((old["sha256"] for old in previous.get("_source", {}).get("workbooks", [])
              if old["file"] == wb["file"]), None) != wb["sha256"]
        for wb in workbooks
    )
    updated_on = (
        dt.datetime.now(dt.timezone.utc).date().isoformat()
        if changed_source else previous.get("_source", {}).get("updatedOn", "2026-10-19")
    )
    return {
        "_source": {"updatedOn": updated_on, "workbooks": workbooks},
        WORKBOOKS[0]: data[WORKBOOKS[0]],
        WORKBOOKS[1]: data[WORKBOOKS[1]],
    }, historical_count, former_count


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--check", action="store_true", help="Fail on stale JSON snapshot")
    mode.add_argument("--write", action="store_true", help="Regenerate the JSON source")
    args = parser.parse_args()
    new, total, old = create_snapshot()
    current = json.loads(SNAPSHOT.read_text(encoding="utf-8"))
    changed = new != current
    meta = new["_source"]
    print(f"Historical results: previous={old}, actual GitHub XLSX={total}")
    for wb in meta["workbooks"]:
        print(f"  {wb['file']}: sha256={wb['sha256']}")
    if args.check:
        if changed:
            print("STALE: backend/data/historical-xlsx.json differs from current GitHub XLSX files.", file=sys.stderr)
            print("Run: python scripts/sync_historical_xlsx.py --write", file=sys.stderr)
            return 1
        print("Historical JSON matches both XLSX workbooks.")
        return 0
    if changed:
        SNAPSHOT.write_text(json.dumps(new, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
        print("UPDATED: backend/data/historical-xlsx.json")
    else:
        print("Already up to date.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
